"""Tests for rag.loaders."""
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from rag.loaders import load_docx, load_pdf, load_text

# Locate the corpus relative to this test file, not the working directory.
RAW_DIR = Path(__file__).resolve().parents[1] / "data" / "raw"


def test_load_text_returns_single_page_list() -> None:
    """A .txt file is one logical "page": a single-element list."""
    pages = load_text(str(RAW_DIR / "handbook.txt"))
    assert isinstance(pages, list)
    assert len(pages) == 1
    assert len(pages[0]) > 0  # content is non-empty


def test_load_text_rejects_wrong_suffix() -> None:
    """Wrong file type is refused before the file is even opened."""
    with pytest.raises(ValueError):
        load_text(str(RAW_DIR / "archive.tar.gz"))


def test_load_text_missing_file_raises() -> None:
    """A missing file must raise, never silently return empty content."""
    with pytest.raises(FileNotFoundError):
        load_text(str(RAW_DIR / "no_such_file.txt"))


def test_load_docx_roundtrip() -> None:
    """Create a tiny .docx on the fly, load it back, verify paragraphs survive.

    Builds the temp dir with ``tempfile.mkdtemp`` rather than pytest's
    ``tmp_path`` fixture — the latter triggers a permission-denied in
    some sandboxed runners.
    """
    import tempfile

    d = Path(tempfile.mkdtemp(prefix="rag_loader_"))
    path = d / "note.docx"
    doc = Document()
    doc.add_paragraph("Annual leave is 15 days.")
    doc.add_paragraph("Remote work allowed twice a week.")
    doc.save(path)

    pages = load_docx(str(path))
    assert len(pages) == 1
    assert "Annual leave is 15 days." in pages[0]
    assert "\n\n" in pages[0]  # paragraph boundary preserved for the chunker


# ---------------------------------------------------------------------------
# PDF loader + OCR fallback
# ---------------------------------------------------------------------------

def test_load_pdf_uses_text_layer_when_available() -> None:
    """When pypdf yields real prose, OCR must NOT be called.

    The check is: if no scan fallback was needed, the OCR engine
    should never be instantiated and pypdfium2 should never be touched.
    """
    import tempfile

    d = Path(tempfile.mkdtemp(prefix="rag_loader_pdf_"))
    path = d / "text.pdf"
    path.write_bytes(b"%PDF-stub")

    class _FakePage:
        def __init__(self, text): self._t = text
        def extract_text(self): return self._t

    class _FakeReader:
        def __init__(self, _path, *a, **kw):
            self.pages = [_FakePage("Real prose here " * 5)]

    with patch("rag.loaders.PdfReader", _FakeReader), \
         patch("rag.loaders._render_page_to_pil") as render, \
         patch("rag.loaders._get_ocr_engine") as engine:
        pages = load_pdf(str(path))
    assert pages == ["Real prose here " * 5]
    render.assert_not_called()
    engine.assert_not_called()


def test_load_pdf_falls_back_to_ocr_on_empty_page() -> None:
    """A page with no text layer must trigger render + OCR.

    We mock the heavy dependencies (pypdfium2 + rapidocr) so the test
    runs in ms, then verify the call chain is exactly: render -> OCR.
    """
    import tempfile

    d = Path(tempfile.mkdtemp(prefix="rag_loader_pdf_ocr_"))
    path = d / "scan.pdf"
    path.write_bytes(b"%PDF-stub")

    class _FakePage:
        def __init__(self, text): self._t = text
        def extract_text(self): return self._t

    class _FakeReader:
        def __init__(self, _path, *a, **kw):
            self.pages = [_FakePage(""), _FakePage("Some prose " * 5)]

    def _fake_render(p, i): return f"PIL-page-{i}"
    def _fake_ocr(img): return f"OCR-of-{img}"

    with patch("rag.loaders.PdfReader", _FakeReader), \
         patch("rag.loaders._render_page_to_pil", side_effect=_fake_render) as render, \
         patch("rag.loaders._ocr_image", side_effect=_fake_ocr) as ocr:
        pages = load_pdf(str(path))

    # Page 0 had empty text -> OCR path. Page 1 had real text -> no OCR.
    assert pages == ["OCR-of-PIL-page-0", "Some prose " * 5]
    assert render.call_count == 1
    assert ocr.call_count == 1
    render.assert_called_once_with(str(path), 0)
