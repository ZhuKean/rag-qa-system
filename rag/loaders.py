"""Document loaders: convert raw files into page-level plain text.

Interface contract: every loader returns ``list[str]`` where one element
corresponds to one "page" of the source (for txt/md the whole file is
treated as a single page). The RAG pipeline then chunks each page
independently with ``rag.chunker.chunk_text``.
"""
from __future__ import annotations

import logging
from functools import lru_cache

from pypdf import PdfReader
from docx import Document

_log = logging.getLogger(__name__)

# Heuristic: if a PDF page's text layer yields fewer than this many
# non-whitespace characters, treat the page as a scan and OCR it.
# 20 chars ≈ a couple of header lines; anything shorter is almost
# certainly an image-only page.
_EMPTY_PAGE_CHARS = 20

# Render scale for OCR. 2.0 = ~144 DPI which is enough for clean
# Chinese print (>=12pt). Bump higher if your scans are small font.
_OCR_RENDER_SCALE = 2.0


# ---------------------------------------------------------------------------
# Plain-text loaders
# ---------------------------------------------------------------------------

def load_text(path: str) -> list[str]:
    """Load .txt/.md: return [whole file text]."""
    if not path.endswith((".txt", ".md")):
        raise ValueError("Path must end with .txt/.md")
    with open(path, "r", encoding="utf-8") as f:
        return [f.read()]


def load_docx(path: str) -> list[str]:
    """Load .docx: return [whole file text]. Using python-docx."""
    if not path.endswith(".docx"):
        raise ValueError("Path must end with .docx")
    document = Document(path)
    text = "\n\n".join(p.text for p in document.paragraphs)
    return [text]


# ---------------------------------------------------------------------------
# PDF with OCR fallback
# ---------------------------------------------------------------------------

@lru_cache(maxsize=1)
def _get_ocr_engine():
    """Lazy-loaded RapidOCR singleton.

    The first call downloads ~30MB of ONNX weights and instantiates
    a detection + recognition pipeline. Cached so we pay the cost
    once per process.
    """
    from rapidocr import RapidOCR  # heavy import — deferred
    return RapidOCR()


def _render_page_to_pil(path: str, page_index: int):
    """Render a single PDF page to a PIL.Image using pypdfium2."""
    import pypdfium2 as pdfium  # heavy import — deferred
    pdf = pdfium.PdfDocument(path)
    page = pdf[page_index]
    bitmap = page.render(scale=_OCR_RENDER_SCALE)
    return bitmap.to_pil()


def _ocr_image(pil_image) -> str:
    """Run OCR on a PIL image and return concatenated text.

    RapidOCR >= 3.x returns a ``RapidOCROutput`` dataclass with three
    tuple fields (``boxes``, ``txts``, ``scores``). Older versions
    returned a ``(result, elapsed)`` tuple. We read ``txts`` directly
    and fall back to an empty string when OCR finds nothing.
    """
    engine = _get_ocr_engine()
    out = engine(pil_image)
    # RapidOCROutput exposes .txts as a tuple[str, ...]; missing on the
    # legacy return shape.
    txts = getattr(out, "txts", None)
    if not txts:
        # Legacy tuple path: ``engine(image)`` returned ``(result, elapsed)``
        try:
            result, _ = out  # type: ignore[misc]
            txts = [line[1] for line in result if line and len(line) >= 2]
        except Exception:
            txts = []
    return "\n".join(txts or [])


def load_pdf(path: str) -> list[str]:
    """Load .pdf: return [page1_text, page2_text, ...].

    Each page is first extracted via pypdf's text layer. If the extracted
    text is too short to be real prose (the page is likely a scan with
    no embedded text), the page is rendered to a bitmap and OCR'd via
    RapidOCR. The caller never has to know whether a page was
    text-extracted or OCR'd — the page-level contract is preserved.
    """
    if not path.endswith(".pdf"):
        raise ValueError("Path must end with .pdf")

    reader = PdfReader(path)
    out: list[str] = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pragma: no cover — defensive
            _log.warning("pypdf failed on page %d of %s: %s", i, path, exc)
            text = ""

        if text.strip() and len(text.strip()) >= _EMPTY_PAGE_CHARS:
            out.append(text)
            continue

        # Page is image-only (or has only a few stray characters).
        # Fall back to OCR. Failure here is non-fatal — we record an
        # empty page so downstream chunking still has the right number
        # of page boundaries.
        _log.info("page %d of %s has no text layer — falling back to OCR", i, path)
        try:
            pil_image = _render_page_to_pil(path, i)
            ocr_text = _ocr_image(pil_image)
        except Exception as exc:  # pragma: no cover — defensive
            _log.warning("OCR failed on page %d of %s: %s", i, path, exc)
            ocr_text = ""
        out.append(ocr_text)

    return out
